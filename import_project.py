#!/usr/bin/env python3
"""
FabriCost PRO style Project Importer Demo
Improved for technical drawing PDFs with element tables.
"""
import os
import sys
import zipfile
import tempfile
import shutil
import subprocess
import re
from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl import load_workbook
import pdfplumber
from pdf2image import convert_from_path
import pytesseract
from PIL import Image
from docx import Document
import ezdxf
import json
from datetime import datetime

SUPPORTED = {
    '.pdf': 'PDF (text + OCR for scanned)',
    '.xlsx': 'Excel XLSX',
    '.xls': 'Excel XLS (via LibreOffice)',
    '.csv': 'CSV',
    '.tsv': 'TSV',
    '.docx': 'Word DOCX',
    '.doc': 'Word DOC (via LibreOffice)',
    '.txt': 'TXT',
    '.rtf': 'RTF (via LibreOffice)',
    '.odt': 'ODT (via LibreOffice)',
    '.ods': 'ODS (via LibreOffice)',
    '.jpg': 'JPG + OCR',
    '.jpeg': 'JPEG + OCR',
    '.png': 'PNG + OCR',
    '.tiff': 'TIFF + OCR',
    '.tif': 'TIF + OCR',
    '.bmp': 'BMP + OCR',
    '.dxf': 'DXF (ezdxf)',
    '.zip': 'ZIP archive of supported',
}

def convert_with_libreoffice(src, target_ext, outdir):
    try:
        cmd = ['soffice', '--headless', '--convert-to', target_ext, '--outdir', outdir, src]
        subprocess.run(cmd, capture_output=True, text=True, timeout=90)
        base = Path(src).stem
        converted = Path(outdir) / f"{base}.{target_ext}"
        return str(converted) if converted.exists() else None
    except Exception:
        return None

def extract_text_from_pdf(path):
    """Efficient extraction: prioritize tables and element-like lines, limit pages if huge."""
    texts = []
    sources = []
    element_records = []
    try:
        with pdfplumber.open(path) as pdf:
            max_pages = min(len(pdf.pages), 50)  # limit for demo performance
            for i, page in enumerate(pdf.pages[:max_pages], 1):
                text = page.extract_text() or ""
                if text.strip():
                    texts.append(text)
                    sources.append(f"PDF puslapis {i}")
                    # Look for element lines: e.g. x28 PL20*500 S355 1 500 0.5 39.2
                    for line in text.split('\n'):
                        if re.search(r'(x\d+|Poz|PLA-|PIS-|PIK-).*?(S\d+|IPE|HEA|HEB|UPN|L\d+|PL\d+)', line, re.I):
                            element_records.append({
                                "raw": line.strip(),
                                "šaltinis": f"PDF puslapis {i}",
                                "tipas": "dokumentas"
                            })
                tables = page.extract_tables()
                for ti, table in enumerate(tables):
                    if table and len(table) > 1:
                        # Try to find header row with Poz/Profilis
                        header_idx = 0
                        for hi, row in enumerate(table[:5]):
                            row_str = ' '.join(str(c or '') for c in row).lower()
                            if 'poz' in row_str or 'profilis' in row_str or 'medžiaga' in row_str:
                                header_idx = hi
                                break
                        headers = [str(c or '').strip() for c in table[header_idx]]
                        for row in table[header_idx+1:]:
                            if any(c for c in row if c and str(c).strip()):
                                rec = {}
                                for h, v in zip(headers, row):
                                    if h and v and str(v).strip():
                                        rec[h] = str(v).strip()
                                if rec:
                                    rec['_source'] = f"PDF puslapis {i} lentelė {ti+1}"
                                    element_records.append(rec)
        if not texts and not element_records:
            # OCR only if no text at all (limit pages)
            images = convert_from_path(path, dpi=150, first_page=1, last_page=min(10, max_pages))
            for i, img in enumerate(images, 1):
                ocr_text = pytesseract.image_to_string(img, lang='eng')
                texts.append(ocr_text)
                sources.append(f"PDF OCR puslapis {i} (patvirtinti)")
        return "\n".join(texts), sources, element_records
    except Exception as e:
        return f"KLAIDA PDF: {e}", ["ERROR"], []

def extract_from_excel(path):
    records = []
    try:
        if path.endswith('.xlsx'):
            wb = load_workbook(path, data_only=True, read_only=True)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    continue
                headers = [str(c or f"col{i}") for i, c in enumerate(rows[0])]
                for idx, row in enumerate(rows[1:], 2):
                    rec = {h: v for h, v in zip(headers, row) if v is not None and str(v).strip()}
                    if rec:
                        rec['_source'] = f"Excel lapas '{sheet_name}' eilutė {idx}"
                        records.append(rec)
            wb.close()
        else:
            with tempfile.TemporaryDirectory() as td:
                conv = convert_with_libreoffice(path, 'xlsx', td)
                if conv:
                    return extract_from_excel(conv)
        return records
    except Exception as e:
        return [{"error": str(e), "_source": path}]

def extract_from_docx(path):
    texts = []
    try:
        doc = Document(path)
        for para in doc.paragraphs:
            if para.text.strip():
                texts.append(para.text)
        for ti, table in enumerate(doc.tables):
            data = [[cell.text for cell in row.cells] for row in table.rows]
            if data:
                texts.append(pd.DataFrame(data[1:], columns=data[0]).to_string())
        return "\n".join(texts), ["DOCX"]
    except Exception as e:
        return f"KLAIDA DOCX: {e}", ["ERROR"]

def extract_from_image(path):
    try:
        img = Image.open(path)
        text = pytesseract.image_to_string(img, lang='eng')
        return text, [f"Vaizdas OCR ({Path(path).name}) - patvirtinti"]
    except Exception as e:
        return f"KLAIDA OCR: {e}", ["ERROR"]

def extract_from_dxf(path):
    texts = []
    try:
        doc = ezdxf.readfile(path)
        for entity in doc.modelspace():
            if entity.dxftype() in ('TEXT', 'MTEXT'):
                t = entity.plain_text() if hasattr(entity, 'plain_text') else str(getattr(entity.dxf, 'text', ''))
                if t.strip():
                    texts.append(t)
        return "\n".join(texts), ["DXF tekstai"]
    except Exception as e:
        return f"KLAIDA DXF: {e}", ["ERROR"]

def process_file(filepath, base_name=None):
    path = Path(filepath)
    ext = path.suffix.lower()
    source_prefix = base_name or path.name
    results = []

    if ext == '.zip':
        with tempfile.TemporaryDirectory() as td:
            with zipfile.ZipFile(path, 'r') as z:
                z.extractall(td)
            for root, _, files in os.walk(td):
                for f in files:
                    fp = os.path.join(root, f)
                    if Path(fp).suffix.lower() in SUPPORTED and not f.endswith('.zip'):
                        results.extend(process_file(fp, f"{source_prefix}/{f}"))
        return results

    if ext == '.pdf':
        text, srcs, element_recs = extract_text_from_pdf(str(path))
        if element_recs:
            for r in element_recs:
                r['_source'] = f"{source_prefix}: {r.get('_source', '')}"
                results.append({"type": "record", "data": r})
        if text.strip():
            results.append({"type": "text", "content": text, "sources": [f"{source_prefix}: {s}" for s in srcs]})
    elif ext in ['.xlsx', '.xls']:
        recs = extract_from_excel(str(path))
        for r in recs:
            r['_source'] = f"{source_prefix}: {r.get('_source', '')}"
            results.append({"type": "record", "data": r})
    elif ext in ['.csv', '.tsv']:
        sep = '\t' if ext == '.tsv' else ','
        try:
            df = pd.read_csv(path, sep=sep)
            for idx, row in df.iterrows():
                rec = {str(k): v for k, v in row.items() if pd.notna(v)}
                rec['_source'] = f"{source_prefix} eilutė {idx+2}"
                results.append({"type": "record", "data": rec})
        except Exception as e:
            results.append({"type": "error", "content": str(e), "source": source_prefix})
    elif ext == '.docx':
        text, srcs = extract_from_docx(str(path))
        results.append({"type": "text", "content": text, "sources": [f"{source_prefix}: {s}" for s in srcs]})
    elif ext in ['.doc', '.rtf', '.odt']:
        with tempfile.TemporaryDirectory() as td:
            conv = convert_with_libreoffice(str(path), 'docx', td)
            if conv:
                text, srcs = extract_from_docx(conv)
                results.append({"type": "text", "content": text, "sources": [f"{source_prefix} (konvertuota)" ]})
            else:
                results.append({"type": "error", "content": "Nepavyko konvertuoti", "source": source_prefix})
    elif ext == '.txt':
        text = path.read_text(encoding='utf-8', errors='ignore')
        results.append({"type": "text", "content": text, "sources": [source_prefix]})
    elif ext == '.ods':
        with tempfile.TemporaryDirectory() as td:
            conv = convert_with_libreoffice(str(path), 'xlsx', td)
            if conv:
                recs = extract_from_excel(conv)
                for r in recs:
                    r['_source'] = f"{source_prefix}: {r.get('_source', '')}"
                    results.append({"type": "record", "data": r})
    elif ext in ['.jpg', '.jpeg', '.png', '.tiff', '.tif', '.bmp']:
        text, srcs = extract_from_image(str(path))
        results.append({"type": "text", "content": text, "sources": [f"{source_prefix}: {s}" for s in srcs], "needs_confirm": True})
    elif ext == '.dxf':
        text, srcs = extract_from_dxf(str(path))
        results.append({"type": "text", "content": text, "sources": [f"{source_prefix}: {s}" for s in srcs]})
    elif ext == '.dwg':
        results.append({"type": "error", "content": "DWG nepalaikomas. Konvertuokite į DXF.", "source": source_prefix})
    else:
        results.append({"type": "error", "content": f"Nepalaikomas formatas: {ext}", "source": source_prefix})
    return results

def parse_for_materials(results):
    materials = []
    project_name = None
    operations = set()
    seen = {}

    op_kw = ['pjovimas', 'suvirinimas', 'gruntavimas', 'dažymas', 'operacija', 'suvirinimo']

    for res in results:
        if res['type'] == 'record':
            data = res['data']
            source = data.get('_source', 'unknown')
            # Map common fields
            pos = data.get('Poz') or data.get('Pozicija') or data.get('poz') or data.get('raw', '')[:20]
            profilis = data.get('Profilis') or data.get('profilis') or data.get('raw')
            medziaga = data.get('Medžiaga') or data.get('medžiaga')
            ilgis = data.get('Ilgis') or data.get('ilgis') or data.get('Ilgis_mm')
            kiekis = data.get('Kiekis') or data.get('kiekis')
            mase = data.get('Svoris') or data.get('Masė') or data.get('mase') or data.get('Masė_kg_m')
            kaina = data.get('Kaina') or data.get('kaina')
            daz = data.get('Daz. plotas') or data.get('Daž. plotas')

            if any([profilis, ilgis, kiekis, mase]) or (isinstance(pos, str) and re.search(r'x\d+|PLA-|PIS-', str(pos))):
                key = f"{profilis}_{ilgis}_{kiekis}"
                entry = {
                    "pozicija": pos,
                    "profilis": profilis or (data.get('raw', '')[:60] if 'raw' in data else None),
                    "medžiaga": medziaga,
                    "ilgis": ilgis,
                    "kiekis": kiekis,
                    "masė": mase,
                    "kaina": kaina,
                    "dažomas_plotas": daz,
                    "šaltinis": source,
                    "tipas": data.get('tipas', 'dokumentas'),
                    "needs_review": data.get('needs_confirm', False)
                }
                if key in seen:
                    entry["needs_review"] = True
                    entry["konfliktas"] = True
                seen[key] = entry
                materials.append(entry)
        elif res['type'] == 'text':
            content = res['content']
            if 'projektas' in content.lower() or 'sandėlio' in content.lower():
                for line in content.split('\n'):
                    if any(k in line.lower() for k in ['projektas', 'sandėlio', 'stogo']):
                        project_name = line.strip()[:120]
                        break
            for op in op_kw:
                if op in content.lower():
                    operations.add(op.capitalize())

    return {
        "projektas": project_name or "Nenustatytas (trūksta aiškaus pavadinimo)",
        "medžiagos": materials,
        "operacijos": list(operations),
        "raw_results_count": len(results)
    }

def generate_html_report(data, output_path):
    # kept for compatibility
    Path(output_path).write_text("<!-- use web UI -->", encoding='utf-8')
    return output_path

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python import_project.py <files...>")
        sys.exit(1)
    all_results = []
    for arg in sys.argv[1:]:
        all_results.extend(process_file(arg))
    parsed = parse_for_materials(all_results)
    print(json.dumps({k: v for k, v in parsed.items()}, ensure_ascii=False, indent=2, default=str))
